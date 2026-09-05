"""Rebuild the bundled Pandoc DOCX reference with justified prose styles."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUTPUT = Path(__file__).parents[1] / "packages/export/src/assets/JustifiedReference.docx"
JUSTIFIED_STYLES = ("Normal", "Body Text", "First Paragraph", "Block Text")


def main() -> None:
    document = Document()
    for name in JUSTIFIED_STYLES:
        if name in document.styles:
            document.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_heading("Markroot justified export reference", level=0)
    document.add_heading("Paragraph styles", level=1)
    document.add_paragraph(
        "This reference document supplies justified prose styles to Pandoc when no explicit "
        "reference document is configured. Its visible content is ignored during export."
    )
    document.add_paragraph("A second paragraph verifies consistent body-text alignment.")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
